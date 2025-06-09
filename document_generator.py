# document_generator.py

import re
import os
import google.generativeai as genai
from io import BytesIO
from docx import Document
from fpdf import FPDF

# Configure Gemini
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
if GEMINI_API_KEY:
    genai.configure(api_key=GEMINI_API_KEY)

# Template descriptions
DOCUMENT_DESCRIPTIONS = {
    "Non-Disclosure Agreement (NDA)": "Protects confidential information shared between two parties.",
    "Employment Agreement": "Outlines the terms and conditions of employment between an employer and employee.",
    "Service Agreement": "Details services to be provided, timelines, fees, and responsibilities.",
    "Freelance Contract": "Defines the working relationship, payment, and scope for freelance work.",
    "Lease Agreement": "Governs the rental of property between a landlord and tenant."
}

# Template file mapping
DOCUMENT_TEMPLATES = {
    "Non-Disclosure Agreement (NDA)": "templates/nda_template.txt",
    "Employment Agreement": "templates/employment_template.txt",
    "Service Agreement": "templates/service_template.txt",
    "Freelance Contract": "templates/freelance_template.txt",
    "Lease Agreement": "templates/lease_template.txt"
}

# Placeholder examples
PLACEHOLDER_EXAMPLES = {
    "start_date": "e.g., January 1, 2025",
    "employer_name": "e.g., ABC Corp.",
    "employee_name": "e.g., John Doe",
    "job_title": "e.g., Software Engineer",
    "supervisor_name": "e.g., Jane Smith",
    "salary": "e.g., $80,000",
    "pay_period": "e.g., annually",
    "work_hours": "e.g., 9 AM - 5 PM",
    "work_location": "e.g., Remote or Office Address",
    "notice_period": "e.g., 2 weeks",
    "clause_confidentiality": "AI-generated",
    "clause_jurisdiction": "AI-generated",
    "project_description": "e.g., Design a company website",
    "fee_amount": "e.g., $1500",
    "payment_milestones": "e.g., 50% upfront, 50% on delivery",
    "due_date": "e.g., March 31, 2025",
    "number_of_revisions": "e.g., 2",
    "clause_ip": "AI-generated",
    "clause_termination": "AI-generated",
    "agreement_date": "e.g., February 1, 2025",
    "client_name": "e.g., XYZ LLC",
    "freelancer_name": "e.g., Alex Freelance",
    "landlord_name": "e.g., Mr. Smith",
    "tenant_name": "e.g., Emily Renter",
    "property_address": "e.g., 123 Main St, NY",
    "lease_start": "e.g., April 1, 2025",
    "lease_end": "e.g., March 31, 2026",
    "monthly_rent": "e.g., $1200",
    "rent_due_date": "e.g., 1st",
    "security_deposit": "e.g., $2400",
    "clause_maintenance": "AI-generated",
    "provider_name": "e.g., Tech Solutions Inc.",
    "service_description": "e.g., IT infrastructure maintenance",
    "service_fee": "e.g., $2000",
    "payment_terms": "e.g., net 30 days",
    "end_date": "e.g., August 1, 2025",
    "clause_dispute": "AI-generated",
    "party_a": "e.g., Innovate Inc.",
    "party_b": "e.g., Recipient Co.",
    "purpose": "e.g., exploring partnership opportunities",
    "duration": "e.g., 2 years",
    "date": "e.g., March 1, 2025"
}

# Utility Functions
def extract_placeholders(template):
    return sorted(set(re.findall(r"{{\s*(\w+)\s*}}", template)))

def load_template(path):
    if os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            return f.read()
    return ""

def generate_clause(prompt):
    try:
        model = genai.GenerativeModel("gemini-1.5-flash")

        response = model.generate_content(prompt)
        return response.text.strip()
    except Exception as e:
        return f"[Error generating clause: {e}]"

def validate_inputs(inputs):
    for key, value in inputs.items():
        if not value.strip():
            return False
    return True

def export_to_docx(text):
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer

def export_to_pdf(text):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", size=12)
    for line in text.splitlines():
        pdf.multi_cell(0, 10, line)
    buffer = BytesIO()
    pdf.output(buffer)
    return buffer
